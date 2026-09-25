import os
import re
from typing import Dict, Any, Optional, List, Union
from pathlib import Path

from backend.services.pdf_inspector import classify_pdf
from backend.services.field_bank import field_bank, CANONICAL_FIELDS, build_header_map, match_header_label
from backend.services.anchors import parse_container_count, parse_gross_weight_kg

class ExtractionResult:
    """
    Standard result object for all document extractions across formats.
    Provides attribute access, dict indexing, and .get() method for complete compatibility.
    """
    def __init__(
        self,
        status: str = "OK",
        reason_code: Optional[str] = None,
        raw_text: str = "",
        fields: Optional[Dict[str, Any]] = None,
        missing_fields: Optional[List[str]] = None,
        audit_trail: Optional[List[Dict[str, Any]]] = None,
        doc_type: str = "UNKNOWN",
        is_unreadable: bool = False,
        is_wrong_doc_type: bool = False,
        confidence: float = 1.0,
        message: Optional[str] = None,
        unresolved_terms: Optional[List[Dict[str, str]]] = None
    ):
        self.status = status
        self.reason_code = reason_code
        self.raw_text = raw_text
        self.fields = fields or {f: None for f in CANONICAL_FIELDS}
        self.missing_fields = missing_fields or []
        self.audit_trail = audit_trail or []
        self.doc_type = doc_type
        self.is_unreadable = is_unreadable
        self.is_wrong_doc_type = is_wrong_doc_type
        self.confidence = confidence
        self.message = message
        self.unresolved_terms = unresolved_terms or []

    # Support dict indexing: result['shipper'], result['status'], etc.
    def __getitem__(self, key: str) -> Any:
        if key in self.fields:
            return self.fields[key]
        if hasattr(self, key):
            return getattr(self, key)
        if key == "review_reason":
            return self.reason_code
        raise KeyError(key)

    def __setitem__(self, key: str, value: Any):
        if key in CANONICAL_FIELDS or key in self.fields:
            self.fields[key] = value
        else:
            setattr(self, key, value)

    def __contains__(self, key: str) -> bool:
        return key in self.fields or hasattr(self, key) or key == "review_reason"

    def get(self, key: str, default: Any = None) -> Any:
        try:
            return self[key]
        except (KeyError, AttributeError):
            return default

    def to_dict(self) -> Dict[str, Any]:
        return {
            "status": self.status,
            "reason_code": self.reason_code,
            "review_reason": self.reason_code,
            "raw_text": self.raw_text,
            "doc_type": self.doc_type,
            "is_unreadable": self.is_unreadable,
            "is_wrong_doc_type": self.is_wrong_doc_type,
            "confidence": self.confidence,
            "missing_fields": self.missing_fields,
            "audit_trail": self.audit_trail,
            "unresolved_terms": self.unresolved_terms,
            **self.fields
        }


class DocumentExtractor:
    """
    Extracts shipping document fields from SI and draft BL documents.
    Routes attachment processing by file extension (.txt, .docx, .xlsx, .pdf).
    """
    # Canonical header patterns (single source of truth: field_bank).
    # Covers plain, parenthetical ("Shipper (Principal or Seller)"),
    # "Consignee (Non-Negotiable)" and Chinese (发货人/收货人/通知方/通知人) variants.
    HEADER_MAP = build_header_map()
    STOP_HEADERS = re.compile(
        r'^(?:document\s*ref|doc\s*ref|document\s*reference|carrier|carrier\s*name|freight\s*status|freight\s*terms|payment\s*terms|payment|terms|date|ocean\s*vessel|vessel\s*name|vessel|export\s*carrier|voy\.\s*no|voyage|voy\b|commodity|description|description\s*of\s*goods|kinds\s*of\s*packages|hs\s*code|booking\s*ref|booking\s*no|booking\s*reference|oc\s*no|freight|bill\s*of\s*lading\s*no|b/l\s*no|b/l\s*number|bl\s*no|order\s*no|bl\s*instruction|bill\s*of\s*lading|container\s*no\b|container\s*no\.|tel\b|fax\b|email\b|p\.?o\.?\s*box|date\b|invoice\s*date|invoice\s*no|inv\s*no|certificate\s*no|country\s*of\s*origin|buyer\b|issuing\s*authority|new\s*no|net\s*weight|tare\s*weight|payment|incoterms|remarks|particulars\s*furnished\s*by\s*shipper)\b',
        re.I
    )

    def extract_attachment(self, path: str, doc_type_hint: str = "SI", email_id: str = "") -> ExtractionResult:
        """
        Single entry point for all attachment types.
        Routes by extension:
        - .txt  -> split(':') key-value logic
        - .docx -> python-docx paragraphs + tables
        - .xlsx -> openpyxl rows
        - .pdf  -> classify_pdf check: text_based (pdfplumber) vs scanned (NEEDS_REVIEW) vs corrupted
        """
        if not os.path.exists(path):
            return ExtractionResult(
                status="NEEDS_REVIEW",
                reason_code="corrupted_file",
                is_unreadable=True,
                message=f"File not found: {path}"
            )

        ext = Path(path).suffix.lower()

        # 1. PDF ROUTE
        if ext == ".pdf":
            try:
                pdf_type = classify_pdf(path)
            except Exception as e:
                # Corrupted or invalid PDF
                return ExtractionResult(
                    status="NEEDS_REVIEW",
                    reason_code="corrupted_file",
                    is_unreadable=True,
                    message=f"Corrupted or invalid PDF file: {str(e)}"
                )

            if pdf_type == "scanned":
                # Scanned image-only PDF — deliberate rule: no OCR, route to human review
                return ExtractionResult(
                    status="NEEDS_REVIEW",
                    reason_code="scanned_not_processed",
                    is_unreadable=True,
                    message="Scanned document (image-only PDF); OCR not attempted by rules-first policy"
                )

            # text_based PDF
            try:
                raw_text = self._extract_pdfplumber_text(path)
                return self.extract_fields(raw_text, doc_type_hint=doc_type_hint, email_id=email_id)
            except Exception as e:
                return ExtractionResult(
                    status="NEEDS_REVIEW",
                    reason_code="corrupted_file",
                    is_unreadable=True,
                    message=f"Could not extract PDF text: {str(e)}"
                )

        # 2. TXT ROUTE
        elif ext == ".txt":
            try:
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    raw_text = f.read()
                return self.extract_fields(raw_text, doc_type_hint=doc_type_hint, email_id=email_id)
            except Exception as e:
                return ExtractionResult(
                    status="NEEDS_REVIEW",
                    reason_code="corrupted_file",
                    is_unreadable=True,
                    message=f"Could not read text file: {str(e)}"
                )

        # 3. DOCX ROUTE
        elif ext == ".docx":
            try:
                import docx
                doc = docx.Document(path)
                lines = []
                for p in doc.paragraphs:
                    t = p.text.strip()
                    if t:
                        lines.append(t)
                for table in doc.tables:
                    for row in table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if len(cells) >= 2:
                            lines.append(f"{cells[0]}: {cells[1]}")
                        elif len(cells) == 1:
                            lines.append(cells[0])
                raw_text = "\n".join(lines)
                return self.extract_fields(raw_text, doc_type_hint=doc_type_hint, email_id=email_id)
            except Exception as e:
                return ExtractionResult(
                    status="NEEDS_REVIEW",
                    reason_code="corrupted_file",
                    is_unreadable=True,
                    message=f"Could not read DOCX file: {str(e)}"
                )

        # 4. XLSX ROUTE
        elif ext == ".xlsx":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(path, data_only=True)
                sheet = wb.active
                lines = []
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if len(cells) >= 2:
                        lines.append(f"{cells[0]}: {cells[1]}")
                    elif len(cells) == 1:
                        lines.append(cells[0])
                raw_text = "\n".join(lines)
                return self.extract_fields(raw_text, doc_type_hint=doc_type_hint, email_id=email_id)
            except Exception as e:
                return ExtractionResult(
                    status="NEEDS_REVIEW",
                    reason_code="corrupted_file",
                    is_unreadable=True,
                    message=f"Could not read XLSX file: {str(e)}"
                )

        # Unsupported or other format
        else:
            try:
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    raw_text = f.read()
                return self.extract_fields(raw_text, doc_type_hint=doc_type_hint, email_id=email_id)
            except Exception as e:
                return ExtractionResult(
                    status="NEEDS_REVIEW",
                    reason_code="corrupted_file",
                    is_unreadable=True,
                    message=f"Unsupported format read error: {str(e)}"
                )

    def _extract_pdfplumber_text(self, path: str) -> str:
        """
        Extracts both plain text and table cells from text-based PDF using pdfplumber.
        Includes optional edgeparse path behind ENABLE_EDGEPARSE feature flag.
        """
        if os.environ.get("ENABLE_EDGEPARSE", "0").lower() in ("1", "true"):
            try:
                # Optional edgeparse path if package is present
                import edgeparse
                return edgeparse.parse(path)
            except Exception:
                pass

        import pdfplumber
        extracted_sections = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                # 1. Plain text
                text = page.extract_text()
                if text:
                    extracted_sections.append(text)

                # 2. Table cells
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            if not row:
                                continue
                            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                            if len(cells) >= 2:
                                extracted_sections.append(f"{cells[0]}: {cells[1]}")
                            elif len(cells) == 1:
                                extracted_sections.append(cells[0])

        return "\n".join(extracted_sections)

    def extract_fields(self, raw_text: str, doc_type_hint: str = "SI", email_id: str = "") -> ExtractionResult:
        """
        Line-by-line field extraction using field_bank and anchors.
        """
        res = ExtractionResult(raw_text=raw_text)

        # 1. Check for unreadable / damaged / corrupted file indicators
        if not raw_text or len(raw_text.strip()) == 0:
            res.status = "NEEDS_REVIEW"
            res.reason_code = "unreadable"
            res.is_unreadable = True
            res.confidence = 0.0
            return res

        raw_lower = raw_text.lower()
        if (
            "[unreadable" in raw_lower or
            "[read_error" in raw_lower or
            "failed to open" in raw_lower or
            "ocr_corrupted" in raw_lower or
            "damaged text" in raw_lower
        ):
            res.status = "NEEDS_REVIEW"
            res.reason_code = "unreadable"
            res.is_unreadable = True
            res.confidence = 0.2
            return res

        # 2. Check for Wrong Document Type
        first_400 = raw_text[:400].upper()
        if "COMMERCIAL INVOICE" in first_400 or "INVOICE NO." in first_400:
            res.status = "NEEDS_REVIEW"
            res.reason_code = "wrong_doc_type"
            res.doc_type = "COMMERCIAL_INVOICE"
            res.is_wrong_doc_type = True
            return res
        elif "PACKING LIST" in first_400:
            res.status = "NEEDS_REVIEW"
            res.reason_code = "wrong_doc_type"
            res.doc_type = "PACKING_LIST"
            res.is_wrong_doc_type = True
            return res
        elif "CERTIFICATE OF ORIGIN" in first_400:
            res.status = "NEEDS_REVIEW"
            res.reason_code = "wrong_doc_type"
            res.doc_type = "CERTIFICATE_OF_ORIGIN"
            res.is_wrong_doc_type = True
            return res
        elif "BILL OF LADING" in first_400:
            res.doc_type = "BILL_OF_LADING"
        elif "SHIPPING INSTRUCTION" in first_400 or "BL INSTRUCTION" in first_400:
            res.doc_type = "SHIPPING_INSTRUCTION"
        else:
            res.doc_type = "BILL_OF_LADING" if doc_type_hint == "BL" else "SHIPPING_INSTRUCTION"

        # 3. Line-by-line parsing
        lines = raw_text.replace('\r', '').splitlines()
        curr_key = None
        curr_val: List[str] = []
        audit_trail: List[Dict[str, Any]] = []
        unresolved_terms: List[Dict[str, str]] = []

        for line in lines:
            line_s = line.strip()
            if not line_s or line_s.startswith("==="):
                continue

            # Check if line contains a key: value separator
            matched_header = False
            if ":" in line_s:
                parts = line_s.split(":", 1)
                cand_label = parts[0].strip()
                rem = parts[1].strip()

                # Check if this label resolves in field_bank
                resolved = field_bank.resolve_label(cand_label)
                if resolved["canonical"] is None:
                    # Deterministic fallback: canonical header patterns cover the
                    # parenthetical / Chinese variants even when they are missing
                    # from the term dictionary or below the fuzzy threshold.
                    header_canonical = match_header_label(cand_label)
                    if header_canonical:
                        resolved = {
                            "canonical": header_canonical,
                            "method": "header_pattern",
                            "confidence": 1.0
                        }
                if resolved["canonical"]:
                    matched_header = True
                    if curr_key:
                        self._set_field_val(res, curr_key, ' '.join(curr_val))
                    curr_key = resolved["canonical"]
                    curr_val = [rem] if rem else []
                    audit_trail.append({
                        "raw_label": cand_label,
                        "canonical": resolved["canonical"],
                        "method": resolved["method"],
                        "confidence": resolved["confidence"]
                    })
                    continue
                elif self.STOP_HEADERS.match(cand_label):
                    matched_header = True
                    if curr_key:
                        self._set_field_val(res, curr_key, ' '.join(curr_val))
                    curr_key = None
                    curr_val = []
                    continue
                elif curr_key in ('shipper', 'consignee', 'notify_party') and cand_label.lower() in ("tel", "phone", "fax", "email", "p.o. box", "p.o.box", "po box", "attn", "attention", "zip"):
                    # Sub-line belonging to address
                    curr_val.append(line_s)
                    continue
                elif len(cand_label) <= 40 and not any(cand_label.lower().startswith(x) for x in ["vessel", "voyage", "booking", "order", "ref", "b/l", "bl ", "attn", "http", "document", "carrier", "freight", "payment", "terms", "date"]):
                    # Unresolved candidate term
                    if rem and rem.upper() not in ("N/A", "NONE", "BLANK", "-"):
                        unresolved_terms.append({"label": cand_label, "value": rem})
                        if email_id:
                            field_bank.record_unresolved_candidate(email_id, cand_label, rem, doc_type=res.doc_type)

            if not matched_header:
                # Check header prefix patterns (e.g. lines like 'Shipper APRIL FINE PAPER TRADING' without colon)
                for pat, canonical in build_header_map():
                    m = pat.match(line_s)
                    if m:
                        if curr_key:
                            self._set_field_val(res, curr_key, ' '.join(curr_val))
                        curr_key = canonical
                        rem = line_s[m.end():].strip()
                        curr_val = [rem] if rem else []
                        audit_trail.append({
                            "raw_label": line_s[:m.end()].strip(),
                            "canonical": canonical,
                            "method": "header_prefix",
                            "confidence": 1.0
                        })
                        matched_header = True
                        break
                if matched_header:
                    continue

            # Continue previous multi-line value
            if curr_key:
                if self.STOP_HEADERS.match(line_s):
                    self._set_field_val(res, curr_key, ' '.join(curr_val))
                    curr_key = None
                    curr_val = []
                else:
                    curr_val.append(line_s)
                    if curr_key in ['port_of_loading', 'port_of_discharge', 'container_count', 'gross_weight_kg']:
                        self._set_field_val(res, curr_key, ' '.join(curr_val))
                        curr_key = None
                        curr_val = []

        if curr_key:
            self._set_field_val(res, curr_key, ' '.join(curr_val))

        # 4. Secondary fallback regex for container count & weight if missed
        if res.fields["container_count"] is None:
            m = re.search(r'(?:total\s*containers|container\s*summary|container\s*count|containers|no\.\s*of\s*containers)[^\d\r\n]*([^\r\n]+)', raw_text, re.I)
            if m:
                res.fields["container_count"] = parse_container_count(m.group(1))

        if res.fields["gross_weight_kg"] is None:
            m = re.search(r'(?:total\s*)?gross\s*(?:weight|wt)[^\d\r\n]*([^\r\n]+)', raw_text, re.I)
            if m:
                res.fields["gross_weight_kg"] = parse_gross_weight_kg(m.group(1))

        # 5. Check for missing required fields
        for f in CANONICAL_FIELDS:
            val = res.fields[f]
            if val is None or (isinstance(val, str) and (val.strip() == "" or "____" in val or val.strip().upper() in ("TBA", "BLANK", "N/A"))):
                res.fields[f] = None
                res.missing_fields.append(f)

        res.audit_trail = audit_trail
        res.unresolved_terms = unresolved_terms

        # Determine review status and reason code
        if res.missing_fields and unresolved_terms:
            res.status = "NEEDS_REVIEW"
            res.reason_code = "term_unresolved"
            res.confidence = 0.75
        elif res.missing_fields:
            res.status = "NEEDS_REVIEW"
            res.reason_code = "missing_value"
            res.confidence = 0.65
        else:
            res.status = "OK"
            res.reason_code = None
            res.confidence = 1.0

        return res

    def _set_field_val(self, res: ExtractionResult, field: str, raw_val: str):
        if not raw_val:
            return

        cleaned = re.sub(r'^(?:\(pol\)|\(pod\)|pol|pod|load port|port of loading|discharge port|port of discharge)\s*[:\-\)]?\s*', '', raw_val, flags=re.I)
        cleaned = cleaned.replace('|', ' ')
        cleaned = re.sub(r'\s+', ' ', cleaned).strip()

        if not cleaned or cleaned.upper() in ["BLANK", "[BLANK]", "[MISSING]", "N/A", "TBA"] or "____" in cleaned:
            res.fields[field] = None
            return

        if field == "container_count":
            res.fields[field] = parse_container_count(cleaned)
        elif field == "gross_weight_kg":
            res.fields[field] = parse_gross_weight_kg(cleaned)
        else:
            res.fields[field] = cleaned


extractor = DocumentExtractor()

def extract_attachment(path: str, doc_type_hint: str = "SI", email_id: str = "") -> ExtractionResult:
    """Convenience functional wrapper around extractor.extract_attachment."""
    return extractor.extract_attachment(path, doc_type_hint=doc_type_hint, email_id=email_id)
