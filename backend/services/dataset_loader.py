import json
import os
import glob
from pathlib import Path
from typing import List, Dict, Any, Optional

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORKSPACE_ROOT = os.path.dirname(BASE_DIR)
TEST_DATA_DIR = os.path.join(WORKSPACE_ROOT, "test data")

class DatasetLoader:
    def __init__(self, data_dir: str = TEST_DATA_DIR):
        self.data_dir = data_dir
        self.inbox_dir = os.path.join(data_dir, "inbox")
        self.attachments_dir = os.path.join(data_dir, "attachments")
        self._emails_cache: Optional[List[Dict[str, Any]]] = None
        self._email_map: Dict[str, Dict[str, Any]] = {}
        self._attachment_text_cache: Dict[str, str] = {}

    def load_inbox(self, force_reload: bool = False) -> List[Dict[str, Any]]:
        if self._emails_cache is not None and not force_reload:
            return self._emails_cache

        emails: List[Dict[str, Any]] = []

        if os.path.exists(self.inbox_dir):
            email_files = sorted(
                glob.glob(os.path.join(self.inbox_dir, "*.json")),
                key=lambda p: (0, int(re_num.group(1))) if (re_num := __import__('re').search(r'email_(\d+)\.json', p)) else (1, str(p))
            )
            for fpath in email_files:
                try:
                    with open(fpath, "r", encoding="utf-8") as f:
                        raw = json.load(f)
                    eid = raw.get("email_id") or raw.get("id") or Path(fpath).stem
                    norm_att = []
                    for att_path in raw.get("attachments", []):
                        if isinstance(att_path, dict):
                            norm_att.append(att_path)
                        else:
                            fname = os.path.basename(att_path)
                            dtype = "SI" if "_si." in fname.lower() else ("BL" if "_bl." in fname.lower() else "OTHER")
                            norm_att.append({
                                "filename": fname,
                                "path": att_path,
                                "doc_type": dtype
                            })

                    norm_email = {
                        "id": eid,
                        "email_id": eid,
                        "sender": raw.get("from") or raw.get("sender", "unknown@shipping.com"),
                        "recipient": raw.get("to") or raw.get("recipient", "ops@maritime-line.com"),
                        "subject": raw.get("subject", ""),
                        "body": raw.get("body", ""),
                        "timestamp": raw.get("timestamp", "2026-01-20T08:30:00Z"),
                        "vessel": raw.get("vessel", self._extract_vessel_hint(raw.get("subject", ""))),
                        "voyage": raw.get("voyage", self._extract_voyage_hint(raw.get("subject", ""))),
                        "company": raw.get("company", self._extract_company_hint(raw.get("subject", ""))),
                        "attachments": norm_att
                    }
                    emails.append(norm_email)
                except Exception as ex:
                    print(f"Error loading {fpath}: {ex}")

        # Fallback to backend/data/inbox.json if test data inbox is not found
        if not emails:
            fallback_file = os.path.join(BASE_DIR, "data", "inbox.json")
            if os.path.exists(fallback_file):
                with open(fallback_file, "r", encoding="utf-8") as f:
                    emails = json.load(f)

        # Prepend dynamically uploaded emails so they appear at the top and survive reload
        uploads_file = os.path.join(WORKSPACE_ROOT, "data", "uploaded_emails.json")
        if os.path.exists(uploads_file):
            try:
                with open(uploads_file, "r", encoding="utf-8") as f:
                    uploaded = json.load(f)
                    if isinstance(uploaded, list):
                        emails = uploaded + emails
            except Exception as ex:
                print(f"[DatasetLoader] Error loading uploaded emails: {ex}")

        self._emails_cache = emails
        self._email_map = {e["id"]: e for e in emails}
        return self._emails_cache

    def register_uploaded_email(self, email_dict: Dict[str, Any]) -> None:
        """Register a new uploaded email document pair persistently."""
        uploads_file = os.path.join(WORKSPACE_ROOT, "data", "uploaded_emails.json")
        os.makedirs(os.path.dirname(uploads_file), exist_ok=True)
        uploaded = []
        if os.path.exists(uploads_file):
            try:
                with open(uploads_file, "r", encoding="utf-8") as f:
                    uploaded = json.load(f)
            except Exception:
                uploaded = []
        # Replace if existing or prepend
        uploaded = [e for e in uploaded if e.get("id") != email_dict.get("id")]
        uploaded.insert(0, email_dict)
        with open(uploads_file, "w", encoding="utf-8") as f:
            json.dump(uploaded, f, indent=2, ensure_ascii=False)

        # Invalidate memory cache so next load reflects update
        self._emails_cache = None
        self._email_map[email_dict["id"]] = email_dict

    def get_email(self, email_id: str) -> Optional[Dict[str, Any]]:
        if self._emails_cache is None:
            self.load_inbox()
        return self._email_map.get(email_id)

    def resolve_attachment_path(self, relative_or_abs_path: str) -> str:
        if os.path.isabs(relative_or_abs_path) and os.path.exists(relative_or_abs_path):
            return relative_or_abs_path

        clean_path = relative_or_abs_path.replace("\\", "/").lstrip("/")
        # Candidates to check
        candidates = [
            os.path.join(self.data_dir, clean_path),
            os.path.join(WORKSPACE_ROOT, clean_path),
            os.path.join(self.attachments_dir, os.path.basename(clean_path)),
            os.path.join(BASE_DIR, "data", clean_path),
            os.path.join(BASE_DIR, "data", "attachments", os.path.basename(clean_path)),
            os.path.join(WORKSPACE_ROOT, "data", "uploads", os.path.basename(clean_path)),
            os.path.join(WORKSPACE_ROOT, "data", "uploads", clean_path),
        ]
        for c in candidates:
            if os.path.exists(c):
                return c
        return candidates[0]

    def read_attachment_bytes(self, relative_or_abs_path: str) -> bytes:
        full_path = self.resolve_attachment_path(relative_or_abs_path)
        if os.path.exists(full_path):
            with open(full_path, "rb") as f:
                return f.read()
        return b""

    def read_attachment_text(self, relative_or_abs_path: str) -> str:
        if relative_or_abs_path in self._attachment_text_cache:
            return self._attachment_text_cache[relative_or_abs_path]

        full_path = self.resolve_attachment_path(relative_or_abs_path)
        if not os.path.exists(full_path):
            return ""

        ext = Path(full_path).suffix.lower()
        text = ""

        try:
            if ext == ".txt":
                with open(full_path, "r", encoding="utf-8", errors="replace") as f:
                    text = f.read()
            elif ext == ".pdf":
                try:
                    from backend.services.pdf_ocr import extract_pdf
                    result = extract_pdf(full_path)
                    text = result["text"] or f"[UNREADABLE_PDF_ERROR: {result['error'] or 'No text found'}]"
                except Exception as ex:
                    text = f"[UNREADABLE_PDF_ERROR: {str(ex)}]"
            elif ext == ".docx":
                try:
                    import docx
                    doc = docx.Document(full_path)
                    paras = [p.text for p in doc.paragraphs if p.text.strip()]
                    for table in doc.tables:
                        for row in table.rows:
                            cells = [c.text.strip() for c in row.cells if c.text.strip()]
                            if len(cells) >= 2:
                                paras.append(f"{cells[0]}: {cells[1]}")
                            elif len(cells) == 1:
                                paras.append(cells[0])
                    text = "\n".join(paras)
                except Exception as ex:
                    text = f"[UNREADABLE_DOCX_ERROR: {str(ex)}]"
            elif ext == ".xlsx":
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(full_path, data_only=True)
                    sheet = wb.active
                    lines = []
                    for row in sheet.iter_rows(values_only=True):
                        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                        if len(cells) >= 2:
                            lines.append(f"{cells[0]}: {cells[1]}")
                        elif len(cells) == 1:
                            lines.append(cells[0])
                    text = "\n".join(lines)
                except Exception as ex:
                    text = f"[UNREADABLE_XLSX_ERROR: {str(ex)}]"
            else:
                with open(full_path, "r", encoding="utf-8", errors="replace") as f:
                    text = f.read()
        except Exception as e:
            text = f"[READ_ERROR: {str(e)}]"

        self._attachment_text_cache[relative_or_abs_path] = text
        return text

    def _extract_vessel_hint(self, subject: str) -> str:
        import re
        m = re.search(r'(?:Vessel|V\.|Draft BL)\s+([A-Z0-9\s]+?)(?:\s+V\.\d+|\s+-\s+|\s*$)', subject, re.I)
        if m:
            return m.group(1).strip()
        return "Commercial Carrier"

    def _extract_voyage_hint(self, subject: str) -> str:
        import re
        m = re.search(r'(V\.[0-9A-Z]+)', subject, re.I)
        return m.group(1) if m else "Standard Voyage"

    def _extract_company_hint(self, subject: str) -> str:
        parts = subject.split("_")
        if len(parts) >= 4:
            return parts[3].strip()
        parts_dash = subject.split("-")
        if len(parts_dash) >= 4:
            return parts_dash[3].strip()
        return "Maritime Shipper"

    def add_uploaded_email(
        self,
        email_dict: Dict[str, Any],
        attachments_content: Optional[Dict[str, Any]] = None
    ) -> Dict[str, Any]:
        """
        Dynamically save a newly uploaded email and its attachment files to disk,
        indexing it into the live inbox dataset so it appears across all views.
        """
        os.makedirs(self.inbox_dir, exist_ok=True)
        os.makedirs(self.attachments_dir, exist_ok=True)

        email_id = email_dict.get("id") or email_dict.get("email_id")
        if not email_id:
            raise ValueError("email_dict must contain 'id' or 'email_id'")

        # Save attachment binary/text files if provided
        if attachments_content:
            for fname, content in attachments_content.items():
                att_path = os.path.join(self.attachments_dir, os.path.basename(fname))
                mode = "wb" if isinstance(content, bytes) else "w"
                kwargs = {} if isinstance(content, bytes) else {"encoding": "utf-8"}
                with open(att_path, mode, **kwargs) as f:
                    f.write(content)
                # clear attachment cache for this file if present
                for k in list(self._attachment_text_cache.keys()):
                    if os.path.basename(fname) in k:
                        self._attachment_text_cache.pop(k, None)

        # Write the email JSON file into inbox_dir
        email_file = os.path.join(self.inbox_dir, f"{email_id}.json")
        with open(email_file, "w", encoding="utf-8") as f:
            json.dump(email_dict, f, indent=2, ensure_ascii=False)

        # Invalidate cache and reload
        self._emails_cache = None
        self._email_map = {}
        self.load_inbox(force_reload=True)
        return self.get_email(email_id) or email_dict

loader = DatasetLoader()
